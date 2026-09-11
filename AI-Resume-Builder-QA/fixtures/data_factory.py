from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from clients.rag import UploadAsset
from fixtures.config import QaSettings


@dataclass(frozen=True)
class RagTestDocument:
    kind: str
    marker: str
    vision_marker: str
    assets: list[UploadAsset]
    expected_file_name: str
    attachment_directory: Path | None = None


class RagDataFactory:
    def __init__(self, root: Path, run_id: str) -> None:
        self._root = root
        self._run_id = run_id
        self._root.mkdir(parents=True, exist_ok=True)

    def create(self, kind: str) -> RagTestDocument:
        run_id = uuid4().hex
        marker = f"QA_RAG_{kind.upper()}_{run_id}"
        vision_marker = "QA_VISION_MARKER"
        if kind == "md":
            image_path = self._create_image(run_id, self._root / "assets")
            return self._create_markdown(marker, vision_marker, image_path)
        image_path = self._create_image(run_id, self._root)
        if kind == "pdf":
            return self._create_pdf(marker, vision_marker, image_path)
        if kind == "docx":
            return self._create_docx(marker, vision_marker, image_path)
        raise ValueError(f"不支持的测试文档类型：{kind}")

    def _create_image(self, run_id: str, directory: Path) -> Path:
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"qa-image-{run_id}.png"
        image = Image.new("RGB", (720, 240), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((12, 12, 708, 228), outline="black", width=3)
        draw.text((40, 90), f"QA VISION SAMPLE {run_id[:12]}", fill="black")
        image.save(path, format="PNG")
        return path

    def _create_markdown(self, marker: str, vision_marker: str, image_path: Path) -> RagTestDocument:
        path = self._root / f"qa-rag-{self._run_id}-{marker.lower()}.md"
        path.write_text(
            f"# QA 知识文档\n\n唯一标记：{marker}\n\n图片标记：{vision_marker}\n\n![QA 图片](assets/{image_path.name})\n",
            encoding="utf-8",
        )
        return RagTestDocument(
            kind="md",
            marker=marker,
            vision_marker=vision_marker,
            assets=[
                UploadAsset(path, "text/markdown", path.name),
                UploadAsset(image_path, "image/png", f"assets/{image_path.name}", role="attachment"),
            ],
            expected_file_name=path.name,
            attachment_directory=image_path.parent,
        )

    def _create_pdf(self, marker: str, vision_marker: str, image_path: Path) -> RagTestDocument:
        path = self._root / f"qa-rag-{self._run_id}-{marker.lower()}.pdf"
        pdf = canvas.Canvas(str(path), pagesize=A4)
        pdf.drawString(72, 780, f"QA marker: {marker}")
        pdf.drawString(72, 755, f"Vision marker: {vision_marker}")
        pdf.drawImage(ImageReader(str(image_path)), 72, 480, width=360, height=120)
        pdf.save()
        return RagTestDocument(
            kind="pdf",
            marker=marker,
            vision_marker=vision_marker,
            assets=[UploadAsset(path, "application/pdf", path.name)],
            expected_file_name=path.name,
        )

    def _create_docx(self, marker: str, vision_marker: str, image_path: Path) -> RagTestDocument:
        path = self._root / f"qa-rag-{self._run_id}-{marker.lower()}.docx"
        document = Document()
        document.add_heading("QA 知识文档", level=1)
        document.add_paragraph(f"唯一标记：{marker}")
        document.add_paragraph(f"图片标记：{vision_marker}")
        document.add_picture(str(image_path), width=Inches(5.0))
        document.save(path)
        return RagTestDocument(
            kind="docx",
            marker=marker,
            vision_marker=vision_marker,
            assets=[
                UploadAsset(
                    path,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    path.name,
                )
            ],
            expected_file_name=path.name,
        )


@pytest.fixture
def rag_data_factory(tmp_path: Path, qa_settings: QaSettings) -> RagDataFactory:
    return RagDataFactory(tmp_path / qa_settings.run_id / "rag-data", qa_settings.run_id)
