# author: jf
import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from docx import Document
from PIL import Image, ImageDraw, PngImagePlugin
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


@dataclass(frozen=True)
class UploadAsset:
    path: Path
    content_type: str
    relative_path: str
    role: str = "document"


@dataclass(frozen=True)
class GeneratedDocument:
    kind: str
    marker: str
    file_name: str
    assets: list[UploadAsset]


class PerformanceDataFactory:
    def __init__(self, run_id: str) -> None:
        self.run_id = run_id
        self.root = Path(__file__).resolve().parents[3] / "reports" / "runtime" / "performance" / run_id / uuid4().hex
        self.root.mkdir(parents=True, exist_ok=True)

    def resume_name(self) -> str:
        return f"qa-resume-{self.run_id}-{uuid4().hex}"

    def resume_data(self, sequence: int = 0) -> dict:
        return {
            "basic": {"name": "QA 脱敏候选人", "title": f"测试开发-{sequence}"},
            "summary": f"仅用于隔离性能测试，run={self.run_id}",
            "skills": ["Python", "API 自动化"],
        }

    def document(self, kind: str = "md", image_count: int = 0) -> GeneratedDocument:
        uid = uuid4().hex
        marker = f"QA_PERF_{self.run_id}_{uid}"
        image_paths = [self._image(uid, index) for index in range(image_count or (1 if kind != "md" else 0))]
        if kind == "md":
            path = self.root / f"qa-rag-{self.run_id}-{uid}.md"
            lines = ["<!-- author: jf -->", "# QA 性能测试知识文档", "", f"唯一标记：{marker}"]
            for index, image_path in enumerate(image_paths, start=1):
                lines.extend(["", f"图片 {index}", f"![QA 图片 {index}](assets/{image_path.name})"])
            path.write_text("\n".join(lines) + "\n", encoding="utf-8")
            assets = [UploadAsset(path, "text/markdown", path.name)]
            assets.extend(UploadAsset(item, "image/png", f"assets/{item.name}", "attachment") for item in image_paths)
        elif kind == "pdf":
            path = self.root / f"qa-rag-{self.run_id}-{uid}.pdf"
            pdf = canvas.Canvas(str(path), pagesize=A4)
            pdf.setAuthor("jf")
            pdf.drawString(72, 780, f"QA marker: {marker}")
            if image_paths:
                pdf.drawImage(ImageReader(str(image_paths[0])), 72, 520, width=360, height=120)
            pdf.save()
            assets = [UploadAsset(path, "application/pdf", path.name)]
        elif kind == "docx":
            path = self.root / f"qa-rag-{self.run_id}-{uid}.docx"
            doc = Document()
            doc.core_properties.author = "jf"
            doc.add_heading("QA 性能测试知识文档", 1)
            doc.add_paragraph(f"唯一标记：{marker}")
            if image_paths:
                doc.add_picture(str(image_paths[0]))
            doc.save(path)
            assets = [UploadAsset(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", path.name)]
        else:
            raise ValueError(f"不支持的文档类型：{kind}")
        return GeneratedDocument(kind, marker, path.name, assets)

    def cleanup(self) -> None:
        runtime_root = (Path(__file__).resolve().parents[3] / "reports" / "runtime" / "performance").resolve()
        target = self.root.resolve()
        if runtime_root not in target.parents or self.run_id not in target.parts:
            raise RuntimeError("本地生成数据目录不满足隔离清理规则")
        shutil.rmtree(target, ignore_errors=True)
        run_root = target.parent
        if run_root.parent == runtime_root and run_root.name == self.run_id and run_root.exists() and not any(run_root.iterdir()):
            run_root.rmdir()

    def _image(self, uid: str, index: int) -> Path:
        directory = self.root / "assets"
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"qa-image-{self.run_id}-{uid}-{index}.png"
        image = Image.new("RGB", (720, 240), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((12, 12, 708, 228), outline="black", width=3)
        draw.text((40, 90), f"QA IMAGE {index} {uid[:10]}", fill="black")
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Author", "jf")
        metadata.add_text("QA", json.dumps({"runId": self.run_id, "index": index}))
        image.save(path, pnginfo=metadata)
        return path
