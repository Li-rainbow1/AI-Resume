from io import BytesIO
import hashlib
import posixpath
from pathlib import Path
from pathlib import PurePosixPath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from app.domain.exceptions.rag_exceptions import FileParseError, UnsupportedFileTypeError, safe_rag_log_value
from app.domain.models.rag_image import DocumentImageCandidate, DocumentImageExtractionResult
from app.domain.models.rag_document import ExtractedDocument


class FileParserAdapter:
    _TEXT_EXTENSIONS = {".txt", ".md"}
    _PDF_EXTENSIONS = {".pdf"}
    _DOCX_EXTENSIONS = {".docx"}
    _IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
    _WORDPROCESSING_NAMESPACE = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    _OFFICE_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def parse(
        self,
        file_bytes: bytes,
        file_name: str,
        content_type: str | None = None,
    ) -> ExtractedDocument:
        safe_name = (file_name or "").strip() or "document"
        extension = Path(safe_name).suffix.lower()
        safe_content_type = (content_type or "").strip() or "application/octet-stream"
        _log_parser(
            "开始解析文档文件",
            file_name=safe_name,
            extension=extension or "none",
            content_type=safe_content_type,
            size_bytes=len(file_bytes),
        )

        if not file_bytes:
            raise FileParseError("上传文件不能为空")

        if extension in self._TEXT_EXTENSIONS:
            content = self._decode_text(file_bytes)
        elif extension in self._PDF_EXTENSIONS:
            content = self._parse_pdf(file_bytes)
        elif extension in self._DOCX_EXTENSIONS:
            content = self._parse_docx(file_bytes)
        else:
            raise UnsupportedFileTypeError(f"暂不支持的文件类型: {extension or safe_content_type}")

        normalized = content.strip()
        if not normalized and extension not in self._PDF_EXTENSIONS | self._DOCX_EXTENSIONS:
            raise FileParseError(f"文件 {safe_name} 未解析到可入库内容")

        if not normalized:
            _log_parser("文档没有文字层，等待图片增强处理", file_name=safe_name, extension=extension)

        source_id = Path(safe_name).stem.strip() or safe_name
        _log_parser("文档解析完成", file_name=safe_name, parsed_chars=len(normalized), source_id=source_id)
        return ExtractedDocument(
            source_id=source_id,
            original_filename=safe_name,
            original_content_type=safe_content_type,
            source_type="document",
            ingest_source="text_document",
            content=normalized,
        )

    def extract_images(
        self,
        file_bytes: bytes,
        file_name: str,
        content_type: str | None = None,
    ) -> DocumentImageExtractionResult:
        """抽取 PDF/DOCX 中可交给视觉模型处理的图片候选。"""
        safe_name = (file_name or "").strip() or "document"
        extension = Path(safe_name).suffix.lower()
        if not file_bytes or extension not in self._PDF_EXTENSIONS | self._DOCX_EXTENSIONS:
            return DocumentImageExtractionResult()

        try:
            if extension in self._PDF_EXTENSIONS:
                candidates, extraction_error = self._extract_pdf_images(file_bytes, safe_name, content_type)
            else:
                candidates, extraction_error = self._extract_docx_images(file_bytes, safe_name, content_type)
        except Exception as exc:
            _log_parser(
                "文档图片抽取失败，正文仍可继续入库",
                file_name=safe_name,
                extension=extension,
                error_type=type(exc).__name__,
            )
            return DocumentImageExtractionResult(
                error_message="复合文档图片抽取失败，可稍后重试图片解析"
            )

        _log_parser(
            "文档图片抽取完成",
            file_name=safe_name,
            extension=extension,
            candidate_count=len(candidates),
            extraction_error=bool(extraction_error),
        )
        return DocumentImageExtractionResult(candidates=candidates, error_message=extraction_error)

    @staticmethod
    def _decode_text(file_bytes: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gbk"):
            try:
                _log_parser("尝试按编码解析文本", encoding=encoding)
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        _log_parser("文本编码兜底解析", encoding="utf-8-replace")
        return file_bytes.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise FileParseError("缺少 pypdf 依赖，无法解析 PDF 文件") from exc

        try:
            reader = PdfReader(BytesIO(file_bytes))
            texts = [(page.extract_text() or "").strip() for page in reader.pages]
            _log_parser("PDF 解析完成", page_count=len(reader.pages))
        except Exception as exc:
            raise FileParseError("PDF 解析失败，请检查文件内容或格式") from exc
        return "\n\n".join(text for text in texts if text)

    @classmethod
    def _extract_pdf_images(
        cls,
        file_bytes: bytes,
        file_name: str,
        content_type: str | None,
    ) -> tuple[list[DocumentImageCandidate], str | None]:
        try:
            from pypdf import PdfReader
        except ImportError:
            _log_parser("缺少 pypdf 依赖，跳过 PDF 图片抽取", file_name=file_name)
            return [], "缺少 PDF 图片抽取依赖，可稍后重试图片解析"

        reader = PdfReader(BytesIO(file_bytes))
        candidates: list[DocumentImageCandidate] = []
        extraction_errors: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            embedded_count = 0
            try:
                page_images = list(page.images)
            except Exception as exc:
                page_images = []
                extraction_errors.append(f"第 {page_number} 页内嵌图片读取失败")
                _log_parser(
                    "PDF 页面内嵌图片读取失败",
                    file_name=file_name,
                    page_number=page_number,
                    error_type=type(exc).__name__,
                )

            for image_index, image in enumerate(page_images, start=1):
                image_bytes = bytes(getattr(image, "data", b"") or b"")
                image_name = str(getattr(image, "name", "") or f"page-{page_number}-image-{image_index}.png")
                candidate = cls._make_candidate(
                    image_bytes=image_bytes,
                    file_name=image_name,
                    content_type=None,
                    source_kind="pdf_embedded",
                    source_locator=f"page/{page_number}/image/{image_index}",
                    image_index=len(candidates),
                    page_number=page_number,
                    paragraph_index=None,
                )
                if candidate is not None:
                    candidates.append(candidate)
                    embedded_count += 1

            if not page_text and embedded_count == 0:
                rendered = cls._render_pdf_page(file_bytes, page_number, file_name)
                if rendered:
                    rendered_candidate = cls._make_candidate(
                        image_bytes=rendered,
                        file_name=f"page-{page_number}.png",
                        content_type="image/png",
                        source_kind="pdf_scanned_page",
                        source_locator=f"page/{page_number}",
                        image_index=len(candidates),
                        page_number=page_number,
                        paragraph_index=None,
                    )
                    if rendered_candidate is not None:
                        candidates.append(rendered_candidate)
                else:
                    extraction_errors.append(f"第 {page_number} 页扫描内容渲染失败")

        error_message = "；".join(extraction_errors[:3]) if extraction_errors else None
        return [candidate for candidate in candidates if candidate is not None], error_message

    @staticmethod
    def _render_pdf_page(file_bytes: bytes, page_number: int, file_name: str) -> bytes | None:
        pdf_document = None
        page = None
        bitmap = None
        image = None
        try:
            import pypdfium2 as pdfium

            pdf_document = pdfium.PdfDocument(file_bytes)
            page = pdf_document[page_number - 1]
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil()
            output = BytesIO()
            image.save(output, format="PNG")
            return output.getvalue()
        except Exception as exc:
            _log_parser(
                "扫描 PDF 页面渲染失败",
                file_name=file_name,
                page_number=page_number,
                error_type=type(exc).__name__,
            )
            return None
        finally:
            for resource in (image, bitmap, page, pdf_document):
                close = getattr(resource, "close", None)
                if callable(close):
                    try:
                        close()
                    except Exception:
                        pass

    @classmethod
    def _extract_docx_images(
        cls,
        file_bytes: bytes,
        file_name: str,
        content_type: str | None,
    ) -> tuple[list[DocumentImageCandidate], str | None]:
        candidates: list[DocumentImageCandidate] = []
        extraction_errors: list[str] = []
        with ZipFile(BytesIO(file_bytes)) as archive:
            content_parts = [
                name
                for name in archive.namelist()
                if (
                    name == "word/document.xml"
                    or (
                        name.startswith(("word/header", "word/footer"))
                        and name.endswith(".xml")
                    )
                )
            ]
            for content_part in sorted(content_parts):
                relationships, relationships_error = cls._read_docx_image_relationships(archive, content_part)
                if relationships_error:
                    extraction_errors.append(relationships_error)
                if not relationships:
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(content_part))
                except (BadZipFile, KeyError, ElementTree.ParseError) as exc:
                    _log_parser(
                        "DOCX 图片所在部件解析失败",
                        file_name=file_name,
                        part_name=content_part,
                        error_type=type(exc).__name__,
                    )
                    extraction_errors.append(f"{content_part} 图片所在部件解析失败")
                    continue

                paragraph_nodes = [
                    node
                    for node in root.iter()
                    if node.tag.rsplit("}", 1)[-1] == "p"
                ]
                for paragraph_index, paragraph in enumerate(paragraph_nodes, start=1):
                    occurrence_index = 0
                    for node in paragraph.iter():
                        if node.tag.rsplit("}", 1)[-1] != "blip":
                            continue
                        occurrence_index += 1
                        relationship_id = node.attrib.get(
                            f"{{{cls._OFFICE_RELATIONSHIP_NAMESPACE}}}embed"
                        )
                        target = relationships.get(relationship_id or "")
                        if not target:
                            continue
                        try:
                            image_bytes = archive.read(target)
                        except KeyError:
                            _log_parser(
                                "DOCX 图片资源不存在",
                                file_name=file_name,
                                part_name=content_part,
                                target=target,
                            )
                            extraction_errors.append(f"{content_part} 图片资源不存在")
                            continue
                        candidate = cls._make_candidate(
                            image_bytes=image_bytes,
                            file_name=Path(target).name,
                            content_type=None,
                            source_kind="docx_embedded",
                            source_locator=(
                                f"{content_part}#paragraph/{paragraph_index}/image/{occurrence_index}"
                            ),
                            image_index=len(candidates),
                            page_number=None,
                            paragraph_index=paragraph_index,
                        )
                        if candidate is not None:
                            candidates.append(candidate)
        error_message = "；".join(extraction_errors[:3]) if extraction_errors else None
        return candidates, error_message

    @classmethod
    def _read_docx_image_relationships(
        cls,
        archive: ZipFile,
        content_part: str,
    ) -> tuple[dict[str, str], str | None]:
        content_path = PurePosixPath(content_part)
        relationships_part = str(
            content_path.parent / "_rels" / f"{content_path.name}.rels"
        )
        try:
            root = ElementTree.fromstring(archive.read(relationships_part))
        except KeyError:
            # 没有关系部件表示该正文/页眉/页脚没有图片，不属于抽取异常。
            return {}, None
        except (BadZipFile, ElementTree.ParseError) as exc:
            return {}, f"{content_part} 图片关系读取失败（{type(exc).__name__}）"

        relationships: dict[str, str] = {}
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] != "Relationship":
                continue
            relationship_id = str(node.attrib.get("Id") or "").strip()
            relationship_type = str(node.attrib.get("Type") or "").strip().lower()
            target_mode = str(node.attrib.get("TargetMode") or "").strip().lower()
            target = str(node.attrib.get("Target") or "").strip()
            if (
                not relationship_id
                or not target
                or target_mode == "external"
                or not relationship_type.endswith("/image")
            ):
                continue
            normalized_target = cls._resolve_docx_target(content_part, target)
            if normalized_target is not None:
                relationships[relationship_id] = normalized_target
        return relationships, None

    @staticmethod
    def _resolve_docx_target(content_part: str, target: str) -> str | None:
        safe_target = target.replace("\\", "/").strip()
        if not safe_target:
            return None
        if safe_target.startswith("/"):
            normalized = posixpath.normpath(safe_target.lstrip("/"))
        else:
            normalized = posixpath.normpath(
                posixpath.join(str(PurePosixPath(content_part).parent), safe_target)
            )
        if (
            normalized in {"", ".", ".."}
            or normalized.startswith("../")
            or not normalized.startswith("word/media/")
            or PurePosixPath(normalized).suffix.lower() not in FileParserAdapter._IMAGE_EXTENSIONS
        ):
            return None
        return normalized

    @classmethod
    def _make_candidate(
        cls,
        *,
        image_bytes: bytes,
        file_name: str,
        content_type: str | None,
        source_kind: str,
        source_locator: str,
        image_index: int,
        page_number: int | None,
        paragraph_index: int | None,
    ) -> DocumentImageCandidate | None:
        if not image_bytes:
            return None
        safe_name = (file_name or "image.png").strip() or "image.png"
        extension = Path(safe_name).suffix.lower()
        resolved_content_type = cls._resolve_image_content_type(safe_name, content_type, image_bytes)
        if extension not in cls._IMAGE_EXTENSIONS and resolved_content_type not in {
            "image/png",
            "image/jpeg",
            "image/webp",
        }:
            return None
        return DocumentImageCandidate(
            image_bytes=image_bytes,
            file_name=safe_name,
            content_type=resolved_content_type,
            source_kind=source_kind,
            source_locator=source_locator,
            image_index=max(0, image_index),
            page_number=page_number,
            paragraph_index=paragraph_index,
            sha256=hashlib.sha256(image_bytes).hexdigest(),
        )

    @staticmethod
    def _resolve_image_content_type(
        file_name: str,
        content_type: str | None,
        image_bytes: bytes = b"",
    ) -> str:
        extension = Path(file_name).suffix.lower()
        if extension == ".png":
            return "image/png"
        if extension in {".jpg", ".jpeg"}:
            return "image/jpeg"
        if extension == ".webp":
            return "image/webp"
        normalized = (content_type or "").split(";", 1)[0].strip().lower()
        if normalized == "image/jpg":
            return "image/jpeg"
        if normalized in {"image/png", "image/jpeg", "image/webp"}:
            return normalized
        # 部分 PDF 内嵌图片没有扩展名，使用文件头兜底识别，避免被误判成不支持格式。
        if image_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
            return "image/png"
        if image_bytes.startswith(b"\xff\xd8\xff"):
            return "image/jpeg"
        if image_bytes[:4] == b"RIFF" and image_bytes[8:12] == b"WEBP":
            return "image/webp"
        return normalized or "application/octet-stream"

    @classmethod
    def _parse_docx(cls, file_bytes: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise FileParseError("缺少 python-docx 依赖，无法解析 DOCX 文件") from exc

        try:
            document = Document(BytesIO(file_bytes))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            if paragraphs:
                _log_parser("DOCX 通过 python-docx 解析成功", paragraph_count=len(paragraphs))
                return "\n\n".join(paragraphs)
        except Exception as exc:
            _log_parser("DOCX 主解析失败，尝试 OOXML 兜底", error_type=type(exc).__name__)
            fallback_text = cls._parse_docx_xml_fallback(file_bytes)
            if fallback_text:
                return fallback_text
            raise FileParseError("DOCX 解析失败，请检查文件内容或格式") from exc

        fallback_text = cls._parse_docx_xml_fallback(file_bytes)
        if fallback_text:
            return fallback_text
        # 有些 DOCX 只有图片没有文字，交给后置图片增强流程继续处理。
        _log_parser("DOCX 未解析到文字层，等待图片增强处理")
        return ""

    @classmethod
    def _parse_docx_xml_fallback(cls, file_bytes: bytes) -> str:
        try:
            with ZipFile(BytesIO(file_bytes)) as archive:
                document_xml = archive.read("word/document.xml")
        except (BadZipFile, KeyError, ValueError):
            return ""

        try:
            root = ElementTree.fromstring(document_xml)
        except ElementTree.ParseError:
            return ""

        paragraphs: list[str] = []
        for paragraph in root.findall(".//w:body//w:p", cls._WORDPROCESSING_NAMESPACE):
            fragments: list[str] = []
            for node in paragraph.iter():
                tag_name = node.tag.rsplit("}", 1)[-1]
                if tag_name == "t" and node.text:
                    fragments.append(node.text)
                elif tag_name == "tab":
                    fragments.append("\t")
                elif tag_name in {"br", "cr"}:
                    fragments.append("\n")
            paragraph_text = "".join(fragments).strip()
            if paragraph_text:
                paragraphs.append(paragraph_text)

        if paragraphs:
            _log_parser("DOCX OOXML 段落兜底解析成功", paragraph_count=len(paragraphs))
            return "\n\n".join(paragraphs)

        text_nodes = [
            (node.text or "").strip()
            for node in root.findall(".//w:t", cls._WORDPROCESSING_NAMESPACE)
            if (node.text or "").strip()
        ]
        if text_nodes:
            _log_parser("DOCX OOXML 文本节点兜底解析成功", node_count=len(text_nodes))
        return "\n".join(text_nodes)


def _log_parser(message: str, **extra: object) -> None:
    parts = [f"[知识库上传][文件解析] {message}"]
    for key, value in extra.items():
        parts.append(f"{key}={safe_rag_log_value(key, value)}")
    print(" ".join(parts), flush=True)
